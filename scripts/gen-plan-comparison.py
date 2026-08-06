#!/usr/bin/env python3
"""Generate the downloadable Health Plans Comparison files (Excel + PDF).

Single source of truth: shared/plan-benefits.ts — the SAME data the Plan
Details page renders. This script parses that file directly (no hand-typed
numbers) so the downloads can never drift from what the site shows. The
row order, section bands, and color cues mirror
client/src/pages/plan-details.tsx.

Outputs (served at https://www.kennion.com/downloads/... via client/public):
  client/public/downloads/Kennion-Health-Plans-Comparison.xlsx
  client/public/downloads/Kennion-Health-Plans-Comparison.pdf

Run after the plan benefits change:
  pip install openpyxl reportlab
  python3 scripts/gen-plan-comparison.py

The buttons on the Plan Details page (client/src/pages/plan-details.tsx)
link to these two static files.
"""
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SRC = REPO / "shared" / "plan-benefits.ts"
OUT_DIR = REPO / "client" / "public" / "downloads"
OUT_DIR.mkdir(parents=True, exist_ok=True)

text = SRC.read_text()

# ── resolve the constants referenced inside the arrays ──────────────────
NC = re.search(r'const NC_AFTER_DED\s*=\s*"([^"]*)"', text).group(1)
ALL_INCLUDE = re.search(r'ALL_PLANS_INCLUDE\s*=\s*\n?\s*"([^"]*)"', text).group(1)


def grab_array(name: str) -> str:
    """Text between `export const NAME ... = [` and its matching `]`."""
    start = text.index(f"export const {name}")
    br = text.index("= [", start) + 2  # skip the `Type[]` annotation's []
    depth, i = 0, br
    while i < len(text):
        c = text[i]
        if c == "[":
            depth += 1
        elif c == "]":
            depth -= 1
            if depth == 0:
                return text[br + 1:i]
        i += 1
    raise ValueError(name)


def parse_objects(body: str):
    """Parse flat `{ key: "val", key: IDENT, ... }` object literals."""
    objs, depth, buf = [], 0, ""
    for c in body:
        if c == "{":
            depth += 1
            if depth == 1:
                buf = ""
                continue
        if c == "}":
            depth -= 1
            if depth == 0:
                objs.append(buf)
                continue
        if depth >= 1:
            buf += c
    out = []
    for raw in objs:
        d = {}
        for km in re.finditer(r'(\w+)\s*:\s*(?:"((?:[^"\\]|\\.)*)"|([A-Za-z_]\w*))', raw):
            key = km.group(1)
            if km.group(2) is not None:
                # UTF-8 safe: only unescape the JS escapes that actually occur
                d[key] = km.group(2).replace('\\"', '"').replace("\\\\", "\\")
            else:
                ident = km.group(3)
                d[key] = NC if ident == "NC_AFTER_DED" else ident
        if d:
            out.append(d)
    return out


MEDICAL = parse_objects(grab_array("MEDICAL_PLAN_DETAILS"))
DENTAL = parse_objects(grab_array("DENTAL_PLAN_DETAILS"))
VISION = parse_objects(grab_array("VISION_PLAN_DETAILS"))
SUPP = parse_objects(grab_array("SUPPLEMENTAL_COVERAGE"))
HOSPITAL = parse_objects(grab_array("HOSPITAL_PLAN_DETAILS"))

print(f"[gen-plan-comparison] medical={len(MEDICAL)} dental={len(DENTAL)} "
      f"vision={len(VISION)} supp={len(SUPP)} hospital={len(HOSPITAL)}",
      file=sys.stderr)
if len(MEDICAL) < 10 or len(DENTAL) < 5 or len(VISION) < 3:
    sys.exit("[gen-plan-comparison] plan counts look wrong — did plan-benefits.ts change shape?")

# Row specs mirror client/src/pages/plan-details.tsx exactly.
# ("section", label) -> full-width band; (label, field) -> data row.
MED_ROWS = [
    ("Deductible", "deductible"), ("Out-of-Pocket Max", "oopMax"),
    ("$0 Preventive Care", "preventiveCare"), ("Benefits App + Concierge", "benefitsApp"),
    ("Visa Card For Expenses", "visaCard"),
    ("section", "Benefits"),
    ("Virtual Primary Care Visits", "virtualPrimary"),
    ("Virtual Mental Health Visits", "virtualMental"),
    ("Virtual Urgent Care", "virtualUrgent"),
    ("Primary Care Office Visits", "primaryCare"),
    ("Specialist Office Visits", "specialist"),
    ("Emergency Room Facility Fee", "er"),
    ("Inpatient Facility Fee", "inpatient"),
    ("Outpatient Facility Fee", "outpatient"),
    ("RX | Generics", "rxGeneric"),
    ("RX | Brand: Preferred", "rxBrandPreferred"),
    ("RX | Brand: Non-preferred", "rxBrandNonPreferred"),
]
DEN_ROWS = [
    ("Deductible", "deductible"), ("Annual Maximum Benefit", "annualMax"),
    ("Lifetime Ortho Maximum", "lifetimeOrthoMax"),
    ("section", "Covered Services"),
    ("Preventative (Plan Pays)", "preventativePct"), ("Basic (Plan Pays)", "basicPct"),
    ("Major (Plan Pays)", "majorPct"), ("Ortho (Plan Pays)", "orthoPct"),
]
VIS_ROWS = [
    ("Exam Copayment", "examCopay"), ("Material Copayment", "materialCopay"),
    ("Exam Every", "examEvery"), ("Lenses Every", "lensesEvery"), ("Frames Every", "framesEvery"),
    ("section", "Allowance"),
    ("Frames", "framesAllowance"), ("Elective Contact Lenses", "electiveContactsAllowance"),
    ("Necessary Contact Lenses", "necessaryContactsAllowance"),
]
HOSP_ROWS = [
    ("section", "Inpatient Hospital"),
    ("Hospital Admission", "inpatientAdmission"),
    ("section", "Other Facility Benefits"),
    ("Emergency Room / Urgent Care Facility", "erUrgentCare"),
    ("section", "Surgeries"),
    ("Inpatient / Outpatient Surgery", "surgery"),
]

# Health (medical) plans only — the downloadable comparison covers the
# medical plans exclusively. Dental / Vision / Supplemental rows still parse
# above (kept in sync for validation) but are not exported.
TABS = [
    ("Medical", MEDICAL, MED_ROWS),
]

# Shared palette (matches the site's navy header + section band + cue colors).
NAVY = "1E3A5F"
BAND = "1F2937"
LABEL_BG = "F1F5F9"
STRIPE = "F8FAFC"
GREEN = "047857"
BLUE = "2563EB"
GREY = "94A3B8"
INK = "0F172A"
FOOTNOTE = ("Benefit descriptions only — not a contract. Rates are quoted separately "
            "and subject to final underwriting. Questions? Hunter Shepherd (hunter@kennion.com).")


def cue_hex(v: str) -> str:
    v = (v or "").strip()
    if v == "Free On App":
        return BLUE
    if v in ("No Charge", "✓"):
        return GREEN
    if v == "—":
        return GREY
    return INK


# ============================ EXCEL ============================
def build_excel(path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    white_bold = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=16, color=NAVY)
    sub_font = Font(italic=True, size=10, color=BLUE)
    label_font = Font(bold=True, size=10, color=INK)
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = Workbook()
    wb.remove(wb.active)

    for tab_name, plans, rows in TABS:
        ws = wb.create_sheet(tab_name)
        ws.sheet_view.showGridLines = False
        ncol = 1 + len(plans)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.cell(1, 1, f"Kennion Health Plans — {tab_name} Comparison").font = title_font
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        c = ws.cell(2, 1, ALL_INCLUDE)
        c.font = sub_font
        c.alignment = left

        hrow = 4
        hc = ws.cell(hrow, 1, "Benefit")
        hc.font = white_bold
        hc.fill = PatternFill("solid", fgColor=NAVY)
        hc.alignment = left
        hc.border = border
        for j, p in enumerate(plans):
            c = ws.cell(hrow, 2 + j, p["name"])
            c.font = white_bold
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = center
            c.border = border
        ws.row_dimensions[hrow].height = 34

        r = hrow + 1
        stripe = False
        for row in rows:
            if row[0] == "section":
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
                c = ws.cell(r, 1, row[1].upper())
                c.font = Font(bold=True, color="FFFFFF", size=9)
                c.fill = PatternFill("solid", fgColor=BAND)
                c.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[r].height = 18
                r += 1
                stripe = False
                continue
            label, field = row
            lc = ws.cell(r, 1, label)
            lc.font = label_font
            lc.alignment = left
            lc.border = border
            lc.fill = PatternFill("solid", fgColor=LABEL_BG)
            for j, p in enumerate(plans):
                v = p.get(field, "")
                c = ws.cell(r, 2 + j, v)
                c.alignment = center
                c.border = border
                c.font = Font(size=10, color=cue_hex(v),
                              bold=v.strip() in ("✓", "Free On App", "No Charge"))
                if stripe:
                    c.fill = PatternFill("solid", fgColor=STRIPE)
            ws.row_dimensions[r].height = 20
            r += 1
            stripe = not stripe

        ws.column_dimensions["A"].width = 30
        for j in range(len(plans)):
            ws.column_dimensions[get_column_letter(2 + j)].width = 15
        ws.freeze_panes = "B5"

        fr = r + 1
        ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=ncol)
        fc = ws.cell(fr, 1, FOOTNOTE)
        fc.font = Font(italic=True, size=9, color="64748B")
        fc.alignment = left

    wb.save(path)
    print(f"[gen-plan-comparison] wrote {path.relative_to(REPO)}", file=sys.stderr)


# ============================ PDF ============================
def build_pdf(path: Path):
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                     Spacer, PageBreak)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT

    navy = colors.HexColor("#" + NAVY)
    band = colors.HexColor("#" + BAND)
    label_bg = colors.HexColor("#" + LABEL_BG)
    stripe_c = colors.HexColor("#" + STRIPE)
    ink = colors.HexColor("#" + INK)
    grid_c = colors.HexColor("#E2E8F0")

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=18, textColor=navy,
                        alignment=TA_LEFT, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9,
                         textColor=colors.HexColor("#" + BLUE), alignment=TA_LEFT,
                         spaceAfter=10, fontName="Helvetica-Oblique")
    cellstyle = ParagraphStyle("cell", parent=styles["Normal"], fontSize=6.5,
                               alignment=1, leading=8, textColor=ink)
    cellstyle_l = ParagraphStyle("cellL", parent=cellstyle, alignment=0,
                                 fontName="Helvetica-Bold")
    head_cell = ParagraphStyle("hcell", parent=cellstyle, textColor=colors.white,
                               fontName="Helvetica-Bold")
    head_cell_l = ParagraphStyle("hcellL", parent=head_cell, alignment=0)
    foot = ParagraphStyle("foot", parent=styles["Normal"], fontSize=7.5,
                          textColor=colors.HexColor("#64748B"),
                          fontName="Helvetica-Oblique", spaceBefore=8)

    def cpar(v):
        st = ParagraphStyle("c", parent=cellstyle, textColor=colors.HexColor("#" + cue_hex(v)),
                            fontName="Helvetica-Bold" if v.strip() in ("✓", "Free On App", "No Charge")
                            else "Helvetica")
        return Paragraph(v.replace("&", "&amp;"), st)

    doc = SimpleDocTemplate(str(path), pagesize=landscape(letter),
                            leftMargin=0.4 * inch, rightMargin=0.4 * inch,
                            topMargin=0.45 * inch, bottomMargin=0.45 * inch)
    page_w = landscape(letter)[0] - 0.8 * inch
    flow = []

    def build_grid(plans, rows):
        label_w = 1.55 * inch
        plan_w = (page_w - label_w) / len(plans)
        data, stys = [], []
        hdr = [Paragraph("BENEFIT", head_cell_l)]
        hdr += [Paragraph(p["name"].replace("&", "&amp;"), head_cell) for p in plans]
        data.append(hdr)
        stys += [("BACKGROUND", (0, 0), (-1, 0), navy),
                 ("TOPPADDING", (0, 0), (-1, 0), 5), ("BOTTOMPADDING", (0, 0), (-1, 0), 5)]
        r = 1
        stripe = False
        for row in rows:
            if row[0] == "section":
                data.append([Paragraph(row[1].upper(), head_cell_l)] + [""] * len(plans))
                stys += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), band),
                         ("TOPPADDING", (0, r), (-1, r), 2), ("BOTTOMPADDING", (0, r), (-1, r), 2)]
                r += 1
                stripe = False
                continue
            label, field = row
            line = [Paragraph(label.replace("&", "&amp;"), cellstyle_l)]
            line += [cpar(p.get(field, "")) for p in plans]
            data.append(line)
            stys.append(("BACKGROUND", (0, r), (0, r), label_bg))
            if stripe:
                stys.append(("BACKGROUND", (1, r), (-1, r), stripe_c))
            r += 1
            stripe = not stripe
        t = Table(data, colWidths=[label_w] + [plan_w] * len(plans), repeatRows=1)
        stys += [
            ("GRID", (0, 0), (-1, -1), 0.4, grid_c),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 1), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 1), (-1, -1), 2.5),
        ]
        t.setStyle(TableStyle(stys))
        return t

    first = True
    for tab_name, plans, rows in TABS:
        if not first:
            flow.append(PageBreak())
        first = False
        flow.append(Paragraph(f"Kennion Health Plans &mdash; {tab_name} Comparison", h1))
        flow.append(Paragraph(ALL_INCLUDE.replace("&", "&amp;"), sub))
        flow.append(build_grid(plans, rows))
        flow.append(Paragraph(FOOTNOTE.replace("&", "&amp;").replace("—", "&mdash;"), foot))

    doc.build(flow)
    print(f"[gen-plan-comparison] wrote {path.relative_to(REPO)}", file=sys.stderr)


if __name__ == "__main__":
    build_excel(OUT_DIR / "Kennion-Health-Plans-Comparison.xlsx")
    build_pdf(OUT_DIR / "Kennion-Health-Plans-Comparison.pdf")
